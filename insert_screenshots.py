#!/usr/bin/env python3
"""Insert PNG screenshots into Word document"""
from docx import Document
from docx.shared import Inches
import os

docx_path = r"D:\\clinic\\Clinic-Platform-User-Manual.docx"
img_dir = r"D:\\clinic"

# Map sections to their screenshots
section_images = [
    (1, "login.png"),
    (2, "admin-users.png"),
    (3, "patients.png"),
    (4, "dashboard.png"),
    (5, "lab-worklist.png"),
    (6, "pharmacy-queue.png"),
    (7, "supply-requisitions.png"),
    (8, "billing-pos.png"),
]

# Open existing document
doc = Document(docx_path)

# Insert images at appropriate sections
img_idx = 0
paragraphs = list(doc.paragraphs)
for i, para in enumerate(paragraphs):
    text = para.text.strip()
    if text.startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.')):
        section_num = int(text.split('.')[0])
        if img_idx < len(section_images) and section_images[img_idx][0] == section_num:
            img_file = os.path.join(img_dir, section_images[img_idx][1])
            if os.path.exists(img_file):
                doc._element.body.insert(i+1, doc.add_picture(img_file, width=Inches(5.5))._element)
                print(f"Inserted {section_images[img_idx][1]} after section {section_num}")
                img_idx += 1

doc.save(docx_path)
print(f"Updated document with {img_idx} images")